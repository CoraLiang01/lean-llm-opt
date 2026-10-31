#!/usr/bin/env python3
"""
export_outputs.py
=================
Put several runs' per-instance outputs side by side, in two forms:

    outputs_4runs/outputs.xlsx     one row per instance, for scanning and sorting
    outputs_4runs/by_instance/     one Markdown file per instance, for reading

    python export_outputs.py                       # the four Large-Scale-OR arms
    python export_outputs.py --arm name=115317 --arm other=130457
    python export_outputs.py --only-disagree       # instances the arms disagree on

Why both formats
----------------
They answer different questions. The workbook is for "which instances went
wrong, and is there a pattern" -- filter, sort, count. The Markdown files are
for "what did it actually produce on question 85", where you need the
formulation, the code and the objective next to each other in full, and a
32k-character spreadsheet cell is the wrong container for a Gurobi model.

The four arms differ in two dimensions that are easy to confuse, so both
outputs name them explicitly rather than by run hash:

    RefData  trimmed | full     which columns the classifier retrieves over
    prompt   pre | post         before/after the candidate-type list was fixed
                                on 2026-08-03 (it had omitted Mixture entirely)

Classification and end-to-end correctness are reported as separate columns
because they move independently: on these runs the arm with the *worst*
classification accuracy had the *best* end-to-end accuracy. Collapsing them
into one "correct" column would hide that.
"""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent

# name -> substring that identifies the run. Order is the column order.
DEFAULT_ARMS = {
    "trimmed_pre":  "20260803-115317",
    "trimmed_post": "20260803-153743",
    "full_pre":     "20260803-130457",
    "full_post":    "20260803-162944",
}

XL_MAX = 32767
ILLEGAL = re.compile(r"[\000-\010\013\014\016-\037]")

# gurobipy's Model.status is an integer. Logged raw it is unreadable, and the
# distinction that matters most here is invisible: 1 (LOADED) means the model
# was built but never solved -- a pipeline failure -- while 3 (INFEASIBLE)
# means it was solved and the formulation was wrong. Both show up as "not 2".
GRB_STATUS = {
    1: "LOADED (built, never solved)", 2: "OPTIMAL", 3: "INFEASIBLE",
    4: "INF_OR_UNBD", 5: "UNBOUNDED", 6: "CUTOFF", 7: "ITERATION_LIMIT",
    8: "NODE_LIMIT", 9: "TIME_LIMIT", 10: "SOLUTION_LIMIT", 11: "INTERRUPTED",
    12: "NUMERIC", 13: "SUBOPTIMAL", 14: "INPROGRESS", 15: "USER_OBJ_LIMIT",
    16: "WORK_LIMIT", 17: "MEM_LIMIT",
}


def status_name(v) -> str:
    if v is None or (isinstance(v, float) and v != v):
        return ""
    try:
        i = int(float(v))
    except (TypeError, ValueError):
        return str(v)
    return f"{i} — {GRB_STATUS.get(i, 'unknown code')}"

FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="D9D9D9")
BAD_FILL = PatternFill("solid", fgColor="FCE4E4")
GOOD_FILL = PatternFill("solid", fgColor="E4F3E4")


def clean(v) -> str:
    """Excel rejects most control characters and truncates past 32767."""
    if v is None or (isinstance(v, float) and v != v):
        return ""
    s = ILLEGAL.sub("", str(v))
    if len(s) > XL_MAX:
        s = s[: XL_MAX - 60] + f"\n\n[... truncated, full text {len(s):,} chars]"
    return s


def find_run(runs_dir: Path, key: str) -> Path:
    hits = sorted(p for p in runs_dir.iterdir()
                  if p.is_dir() and key in p.name)
    if not hits:
        raise SystemExit(f"no run directory matching {key!r} under {runs_dir}")
    return hits[-1]


def load_arm(runs_dir: Path, key: str, scored: pd.DataFrame) -> pd.DataFrame:
    d = find_run(runs_dir, key)
    rows = [json.loads(l) for l in (d / "instances.jsonl").open(encoding="utf-8")]
    df = pd.DataFrame(rows)
    keep = [c for c in ("instance_id", "query", "gold_type", "pred_type",
                        "status", "model_output", "code_output",
                        "gold_objective", "size_class", "dataset_address")
            if c in df.columns]
    df = df[keep].copy()
    df["run_id"] = d.name

    s = scored[scored.run_id == d.name]
    if s.empty:                       # scored_all keys on the full run_id
        s = scored[scored.run_id.str.contains(key, na=False)]
    cols = [c for c in ("instance_id", "objective", "correct_optimal",
                        "solver_status", "solve_error") if c in s.columns]
    return df.merge(s[cols].drop_duplicates("instance_id"),
                    on="instance_id", how="left")


# --------------------------------------------------------------------------- #
def build_workbook(arms: dict[str, pd.DataFrame], out: Path,
                   canon) -> None:
    import openpyxl

    names = list(arms)
    base = arms[names[0]][["instance_id", "size_class", "gold_type",
                           "gold_objective", "query"]].copy()

    cls = base[["instance_id", "size_class", "gold_type"]].copy()
    e2e = base[["instance_id", "size_class", "gold_type",
                "gold_objective"]].copy()
    for n in names:
        a = arms[n].set_index("instance_id")
        cls[f"pred__{n}"] = cls.instance_id.map(a.pred_type)
        cls[f"clsok__{n}"] = [
            int(canon(g) is not None and canon(g) == canon(p))
            for g, p in zip(cls.gold_type, cls[f"pred__{n}"])]
        e2e[f"obj__{n}"] = e2e.instance_id.map(a.objective)
        e2e[f"ok__{n}"] = e2e.instance_id.map(
            a.correct_optimal).fillna(0).astype(int)
        e2e[f"status__{n}"] = e2e.instance_id.map(
            a.get("solver_status")).map(status_name)

    cls["n_arms_correct"] = cls[[f"clsok__{n}" for n in names]].sum(axis=1)
    e2e["n_arms_correct"] = e2e[[f"ok__{n}" for n in names]].sum(axis=1)

    with pd.ExcelWriter(out, engine="openpyxl") as xl:
        pd.DataFrame({"placeholder": []}).to_excel(xl, sheet_name="Summary", index=False)
        cls.to_excel(xl, sheet_name="Classification", index=False)
        e2e.to_excel(xl, sheet_name="EndToEnd", index=False)
        for n in names:
            a = arms[n]
            t = a[["instance_id", "gold_type", "pred_type", "objective",
                   "gold_objective", "correct_optimal", "solver_status",
                   "solve_error", "model_output", "code_output"]].copy()
            t["solver_status"] = t.solver_status.map(status_name)
            for c in ("model_output", "code_output"):
                t[c] = t[c].map(clean)
            t.to_excel(xl, sheet_name=f"out_{n}"[:31], index=False)

    wb = openpyxl.load_workbook(out)
    _write_summary(wb, names, arms, len(base))
    for ws in wb.worksheets:
        _style(ws)
    wb.save(out)


def _write_summary(wb, names, arms, n_inst):
    ws = wb["Summary"]
    ws.delete_cols(1)
    rows = [
        ["LEAN-LLM-OPT / Large-Scale-OR — four arms", "", "", "", ""],
        [f"{n_inst} instances per arm. gpt-4.1, temperature 0.", "", "", "", ""],
        ["", "", "", "", ""],
        ["arm", "RefData columns", "classifier prompt",
         "classification correct", "end-to-end correct"],
    ]
    for r in rows:
        ws.append(r)
    # Counts are written as values, not SUM() formulas. LibreOffice cannot
    # recalculate this workbook in reasonable time -- 400 rows carrying whole
    # Gurobi programs -- and an un-recalculated formula reads back as blank in
    # every previewer and in pandas. A number that is visibly there beats a
    # formula that is only correct once Excel has opened it.
    for n in names:
        cls_col = _col(wb, "Classification", "clsok__" + n)
        e2e_col = _col(wb, "EndToEnd", "ok__" + n)
        cls_n = sum(wb["Classification"][f"{cls_col}{r}"].value or 0
                    for r in range(2, n_inst + 2))
        e2e_n = sum(wb["EndToEnd"][f"{e2e_col}{r}"].value or 0
                    for r in range(2, n_inst + 2))
        ws.append([
            n,
            "trimmed (prompt, Data_address, New Problem Type)"
            if n.startswith("trimmed") else
            "full (adds Label, Label_Code, Related)",
            "pre-fix (no Mixture in the candidate list)"
            if n.endswith("_pre") else "post-fix (Mixture added 2026-08-03)",
            f"{int(cls_n)} / {n_inst}   ({100 * cls_n / n_inst:.1f}%)",
            f"{int(e2e_n)} / {n_inst}   ({100 * e2e_n / n_inst:.1f}%)",
        ])
    ws.append([])
    ws.append([f"Counts are column sums of Classification!clsok__* and "
               f"EndToEnd!ok__* (rows 2:{n_inst + 1}), computed when this file "
               f"was written. Re-run export_outputs.py to refresh them.",
               "", "", "", ""])
    ws.append(["Noise floor: 19 of 101 instances change outcome between two "
               "runs of an identical configuration; net differences have "
               "sd ~4.4 instances. Treat a gap under ~9 instances as noise.",
               "", "", "", ""])
    ws.append(["Source: runs/*/instances.jsonl (model_output, code_output) and "
               "scored_all.csv (objective, correct_optimal), "
               "written by export_outputs.py.", "", "", "", ""])
    for r in (1, 2):
        ws.cell(r, 1).font = Font(FONT, bold=(r == 1), size=13 if r == 1 else 10)
    for c in range(1, 6):
        ws.cell(4, c).font = Font(FONT, bold=True)
        ws.cell(4, c).fill = HDR_FILL
    for w, c in zip((16, 46, 42, 22, 20), "ABCDE"):
        ws.column_dimensions[c].width = w


def _col(wb, sheet, header) -> str:
    ws = wb[sheet]
    for c in range(1, ws.max_column + 1):
        if ws.cell(1, c).value == header:
            return get_column_letter(c)
    raise KeyError(f"{header} not in {sheet}")


def _style(ws):
    if ws.title == "Summary":
        return
    ws.freeze_panes = "B2"
    for c in range(1, ws.max_column + 1):
        h = ws.cell(1, c)
        h.font = Font(FONT, bold=True)
        h.fill = HDR_FILL
        h.alignment = Alignment(wrap_text=True, vertical="top")
        name = str(h.value or "")
        wide = name in ("model_output", "code_output", "query")
        ws.column_dimensions[get_column_letter(c)].width = 60 if wide else 15
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(FONT)
            cell.alignment = Alignment(wrap_text=False, vertical="top")
    # colour the correctness columns: the eye should find the zeros
    for c in range(1, ws.max_column + 1):
        if not str(ws.cell(1, c).value or "").startswith(
                ("clsok__", "ok__", "correct_optimal")):
            continue
        for r in range(2, ws.max_row + 1):
            v = ws.cell(r, c).value
            ws.cell(r, c).fill = GOOD_FILL if v in (1, 1.0) else BAD_FILL


# --------------------------------------------------------------------------- #
MD = """# Instance {iid}   (gold type: {gold}, size {size})

**Gold objective:** `{gobj}`

## Query

{query}

**Data:** {addr}

---

{arms}
"""

ARM_MD = """## {name}

| | |
|---|---|
| classified as | **{pred}** {clsmark} |
| objective | `{obj}` {okmark} |
| solver status | {status} |
| solve error | {err} |

<details><summary>Model (LaTeX)</summary>

```
{model}
```

</details>

<details><summary>Gurobi code</summary>

```python
{code}
```

</details>

"""


def _t(v, empty: str = "—") -> str:
    """pandas hands back float('nan') for every missing cell; `nan` on the page
    reads as a value the model produced, which is the opposite of the truth."""
    if v is None or (isinstance(v, float) and v != v):
        return empty
    s = str(v).strip()
    return s if s and s.lower() != "nan" else empty


def write_markdown(arms: dict[str, pd.DataFrame], out_dir: Path, canon) -> int:
    out_dir.mkdir(parents=True, exist_ok=True)
    names = list(arms)
    idx = arms[names[0]].set_index("instance_id")
    n = 0
    for iid, base in idx.iterrows():
        blocks = []
        for nm in names:
            a = arms[nm].set_index("instance_id")
            if iid not in a.index:
                continue
            r = a.loc[iid]
            ok = float(r.get("correct_optimal") or 0) > 0
            cls_ok = (canon(base.gold_type) is not None
                      and canon(base.gold_type) == canon(r.pred_type))
            blocks.append(ARM_MD.format(
                name=nm, pred=_t(r.pred_type),
                clsmark="✓" if cls_ok else f"✗ (gold {base.gold_type})",
                obj=_t(r.get("objective"), "no objective"),
                okmark="✓" if ok else "✗",
                status=status_name(r.get("solver_status"))
                       or _t(r.get("status")),
                err=_t(r.get("solve_error"), "—"),
                model=str(r.get("model_output") or "").strip(),
                code=str(r.get("code_output") or "").strip()))
        (out_dir / f"instance_{int(iid):03d}.md").write_text(
            MD.format(iid=iid, gold=base.gold_type, size=base.size_class,
                      gobj=base.gold_objective,
                      query=str(base.query).strip(),
                      addr=str(base.dataset_address or "").replace("\n", " · "),
                      arms="\n".join(blocks)), encoding="utf-8")
        n += 1
    return n


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--runs", default="runs")
    ap.add_argument("--scored", default="scored_all.csv")
    ap.add_argument("--config", default="exp_config.yaml")
    ap.add_argument("--out", default="outputs_4runs")
    ap.add_argument("--arm", action="append", metavar="NAME=SUBSTRING",
                    help="repeatable; overrides the default four arms")
    ap.add_argument("--only-disagree", action="store_true",
                    help="Markdown only for instances where the arms differ")
    args = ap.parse_args()

    import sys
    sys.path.insert(0, str(HERE))
    import yaml
    from aggregate_runs import _canon_map, canonical_type
    cls_cfg = (yaml.safe_load(Path(args.config).read_text(encoding="utf-8"))
               .get("common", {}).get("classification", {}) or {})
    cmap = _canon_map(cls_cfg)
    others = bool(cls_cfg.get("others_prefix_counts_as_others", True))
    canon = lambda v: canonical_type(v, cmap, others)     # noqa: E731

    spec = DEFAULT_ARMS
    if args.arm:
        spec = dict(a.split("=", 1) for a in args.arm)

    scored = pd.read_csv(args.scored)
    arms = {n: load_arm(Path(args.runs), k, scored) for n, k in spec.items()}
    for n, df in arms.items():
        print(f"  {n:14s} {len(df):3d} instances  {df.run_id.iloc[0]}")

    out_dir = Path(args.out)
    out_dir.mkdir(parents=True, exist_ok=True)
    xlsx = out_dir / "outputs.xlsx"
    build_workbook(arms, xlsx, canon)
    print(f"\nworkbook -> {xlsx}")

    if args.only_disagree:
        first = arms[list(arms)[0]].set_index("instance_id")
        keep = []
        for iid in first.index:
            vals = {float(arms[n].set_index("instance_id")
                          .loc[iid, "correct_optimal"] or 0) > 0 for n in arms}
            if len(vals) > 1:
                keep.append(iid)
        arms = {n: df[df.instance_id.isin(keep)] for n, df in arms.items()}
        print(f"  restricted to {len(keep)} instances the arms disagree on")

    md_dir = out_dir / "by_instance"
    n = write_markdown(arms, md_dir, canon)
    print(f"markdown -> {md_dir}/  ({n} files)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
